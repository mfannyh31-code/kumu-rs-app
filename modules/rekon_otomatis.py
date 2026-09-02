import streamlit as st
import pandas as pd
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import streamlit.components.v1 as components
from datetime import datetime, date
from db import get_db, format_rupiah, render_header

def format_angka(val):
    try: return f"{int(float(val)):,}".replace(",", ".")
    except: return "0"

def render_page():
    render_header("📑 Rekon & Monitoring Data Bendahara", "Analisis pencocokan rekening koran, rekapitulasi unit, tindakan, rentang tanggal kustom, pencocokan selisih bank, dan cetak laporan.")

    st.markdown("""
        <style>
        .metric-card { background: #F8FAFC; padding: 16px; border-radius: 10px; border: 1px solid #E2E8F0; text-align: center; }
        .tbl-report { width: 100%; border-collapse: collapse; margin-top: 15px; font-size: 13px; border: 1px solid #CBD5E1; }
        .tbl-report th { background: #028090; color: white; padding: 10px; border: 1px solid #026b78; text-align: left; }
        .tbl-report td { padding: 9px; border: 1px solid #CBD5E1; }
        .total-row { font-weight: 800; background: #E2E8F0; border-top: 2px solid #94A3B8; }
        .subtotal-unit-row { font-weight: 700; background: #F1F5F9; color: #0F172A; }
        .compact-filter-box { background: #FFFFFF; padding: 20px; border-radius: 10px; border: 1px solid #E2E8F0; margin-bottom: 20px; }
        .bank-match-box { background: #F0FDF4; padding: 20px; border-radius: 10px; border: 1px solid #BBF7D0; margin-bottom: 20px; }
        @media print {
            .compact-filter-box, .bank-match-box, stSidebar, header, button { display: none !important; }
        }
        </style>
    """, unsafe_allow_html=True)

    conn = get_db()
    
    try:
        # Ambil Kategori Layanan Utama (Service Categories) untuk filter hirarki
        service_cats_df = pd.read_sql_query("SELECT id, name FROM service_categories ORDER BY name ASC", conn)
        scats_dict = dict(zip(service_cats_df['name'], service_cats_df['id']))
        list_service_cats = ["Semua Kategori Layanan"] + list(scats_dict.keys())
    except:
        scats_dict = {}
        list_service_cats = ["Semua Kategori Layanan"]

    try:
        categories_df = pd.read_sql_query("SELECT id, service_category_id, name FROM categories ORDER BY name ASC", conn)
        cats_dict = dict(zip(categories_df['name'], categories_df['id']))
        list_units = ["Semua Unit"] + list(cats_dict.keys())
    except:
        cats_dict = {}
        list_units = ["Semua Unit"]

    try:
        df_users = pd.read_sql_query("SELECT DISTINCT cashier_username FROM transactions WHERE cashier_username IS NOT NULL", conn)
        list_users = ["Semua Kasir"] + [str(u).upper() for u in df_users['cashier_username'].tolist() if str(u).strip() != ""]
    except:
        list_users = ["Semua Kasir"]

    # --- FILTER SECTION ---
    st.markdown('<div class="compact-filter-box">', unsafe_allow_html=True)
    
    c1, c2, c3, c4 = st.columns([1.2, 1.2, 1.3, 1.3])
    with c1:
        tipe_laporan = st.selectbox("Jenis Periode", ["Harian", "Bulanan", "Tahunan", "Rentang Tanggal (Custom)"])
    
    months_dict = {
        "Januari": "01", "Februari": "02", "Maret": "03", "April": "04", 
        "Mei": "05", "Juni": "06", "Juli": "07", "Agustus": "08", 
        "September": "09", "Oktober": "10", "November": "11", "Desember": "12"
    }
    years_list = [str(y) for y in range(2024, 2032)]
    current_year_str = str(datetime.today().year)

    if tipe_laporan == "Harian":
        with c2:
            filter_val = st.date_input("Pilih Tanggal", datetime.today())
        date_mask_sql = f"{str(filter_val)}%"
        date_filter_mode = "like"
    elif tipe_laporan == "Rentang Tanggal (Custom)":
        with c2:
            start_d = st.date_input("Dari Tanggal", datetime.today())
        with c3:
            end_d = st.date_input("Sampai Tanggal", datetime.today())
        date_filter_mode = "range"
    elif tipe_laporan == "Bulanan":
        with c2:
            sel_month_name = st.selectbox("Pilih Bulan", list(months_dict.keys()), index=datetime.today().month - 1)
        with c3:
            sel_year_m = st.selectbox("Pilih Tahun (Bulan)", years_list, index=years_list.index(current_year_str) if current_year_str in years_list else 0)
        m_num = months_dict[sel_month_name]
        date_mask_sql = f"{sel_year_m}-{m_num}%"
        date_filter_mode = "like"
    else:
        with c2:
            sel_year = st.selectbox("Pilih Tahun", years_list, index=years_list.index(current_year_str) if current_year_str in years_list else 0)
        date_mask_sql = f"{sel_year}%"
        date_filter_mode = "like"

    f_u0, f_u1, f_u2, f_u3, f_u4, f_u5 = st.columns([1.3, 1.2, 1.2, 1.2, 1.2, 1.5])
    with f_u0:
        selected_service_cat = st.selectbox("Filter Kategori Utama", list_service_cats)
    with f_u1:
        selected_unit = st.selectbox("Filter Unit", list_units)

    try:
        if selected_unit != "Semua Unit" and selected_unit in cats_dict:
            cat_id = cats_dict[selected_unit]
            actions_df = pd.read_sql_query("SELECT name FROM actions WHERE category_id = ? ORDER BY name ASC", conn, params=[cat_id])
        else:
            actions_df = pd.read_sql_query("SELECT name FROM actions ORDER BY name ASC", conn)
        
        list_actions = ["Semua Tindakan"] + actions_df['name'].tolist()
    except:
        list_actions = ["Semua Tindakan"]

    with f_u2:
        selected_action = st.selectbox("Filter Tindakan", list_actions)
    with f_u3:
        selected_kasir_filter = st.selectbox("Filter Kasir", list_users)
    with f_u4:
        selected_shift_filter = st.selectbox("Filter Shift", ["Semua Shift", "Pagi", "Sore", "Malam"])
    with f_u5:
        search_keyword = st.text_input("🔍 Cari No. Kertas", placeholder="Ketik kata kunci...")

    st.markdown('</div>', unsafe_allow_html=True)

    # --- QUERY DATABASE BERDASARKAN FILTER WAKTU & HIRARKI BARU ---
    if date_filter_mode == "range":
        q_items = """
            SELECT i.category_name, i.action_name, i.price, i.qty, i.discount, i.subtotal, t.receipt_date, t.receipt_no, t.cashier_username, t.shift,
                   COALESCE(sc.name, 'Lainnya') as service_category_name
            FROM transaction_items i
            JOIN transactions t ON i.receipt_no = t.receipt_no
            LEFT JOIN categories c ON i.category_name = c.name
            LEFT JOIN service_categories sc ON c.service_category_id = sc.id
            WHERE SUBSTR(t.receipt_date, 1, 10) BETWEEN ? AND ?
        """
        params_items = [str(start_d), str(end_d)]
    else:
        q_items = """
            SELECT i.category_name, i.action_name, i.price, i.qty, i.discount, i.subtotal, t.receipt_date, t.receipt_no, t.cashier_username, t.shift,
                   COALESCE(sc.name, 'Lainnya') as service_category_name
            FROM transaction_items i
            JOIN transactions t ON i.receipt_no = t.receipt_no
            LEFT JOIN categories c ON i.category_name = c.name
            LEFT JOIN service_categories sc ON c.service_category_id = sc.id
            WHERE t.receipt_date LIKE ?
        """
        params_items = [date_mask_sql]

    if selected_service_cat != "Semua Kategori Layanan":
        q_items += " AND sc.name = ?"
        params_items.append(selected_service_cat)
    if selected_unit != "Semua Unit":
        q_items += " AND i.category_name = ?"
        params_items.append(selected_unit)
    if selected_action != "Semua Tindakan":
        q_items += " AND i.action_name = ?"
        params_items.append(selected_action)
    if selected_kasir_filter != "Semua Kasir":
        q_items += " AND UPPER(TRIM(t.cashier_username)) = UPPER(TRIM(?))"
        params_items.append(selected_kasir_filter)
    if selected_shift_filter != "Semua Shift":
        q_items += " AND UPPER(TRIM(t.shift)) = UPPER(TRIM(?))"
        params_items.append(selected_shift_filter)
        
    df_items = pd.read_sql_query(q_items, conn, params=params_items)
    valid_receipts = df_items['receipt_no'].unique().tolist() if not df_items.empty else []

    if date_filter_mode == "range":
        df_tx = pd.read_sql_query("SELECT * FROM transactions WHERE SUBSTR(receipt_date, 1, 10) BETWEEN ? AND ?", conn, params=[str(start_d), str(end_d)])
    else:
        if selected_service_cat == "Semua Kategori Layanan" and selected_unit == "Semua Unit" and selected_action == "Semua Tindakan" and selected_kasir_filter == "Semua Kasir" and selected_shift_filter == "Semua Shift":
            df_tx = pd.read_sql_query("SELECT * FROM transactions WHERE receipt_date LIKE ?", conn, params=[date_mask_sql])
        else:
            if valid_receipts:
                placeholders = ','.join(['?'] * len(valid_receipts))
                df_tx = pd.read_sql_query(f"SELECT * FROM transactions WHERE receipt_date LIKE ? AND receipt_no IN ({placeholders})", conn, params=[date_mask_sql] + valid_receipts)
            else:
                df_tx = pd.DataFrame()

    try:
        if date_filter_mode == "range":
            piu_query = "SELECT * FROM receivables_payments WHERE SUBSTR(pay_date, 1, 10) BETWEEN ? AND ?"
            piu_params = [str(start_d), str(end_d)]
        else:
            piu_query = "SELECT * FROM receivables_payments WHERE pay_date LIKE ?"
            piu_params = [date_mask_sql]

        if selected_kasir_filter != "Semua Kasir":
            piu_query += " AND UPPER(TRIM(input_by)) = UPPER(TRIM(?))"
            piu_params.append(selected_kasir_filter)
        if selected_shift_filter != "Semua Shift":
            piu_query += " AND UPPER(TRIM(shift)) = UPPER(TRIM(?))"
            piu_params.append(selected_shift_filter)
        df_piu = pd.read_sql_query(piu_query, conn, params=piu_params)
    except:
        df_piu = pd.DataFrame()
    
    conn.close()

    if search_keyword.strip() and not df_items.empty:
        kw = search_keyword.strip().lower()
        df_items = df_items[
            df_items['receipt_no'].astype(str).str.lower().str.contains(kw) |
            df_items['action_name'].astype(str).str.lower().str.contains(kw) |
            df_items['category_name'].astype(str).str.lower().str.contains(kw) |
            df_items['service_category_name'].astype(str).str.lower().str.contains(kw)
        ]

    def is_tunai(method_str):
        if not method_str: return False
        return "TUNAI" in str(method_str).upper().strip()

    piu_tunai = df_piu[df_piu['method'].apply(is_tunai)]['amount'].sum() if not df_piu.empty and 'amount' in df_piu.columns else 0.0
    piu_transfer = df_piu[df_piu['method'].str.upper().str.contains('TRANSFER', na=False)]['amount'].sum() if not df_piu.empty and 'amount' in df_piu.columns else 0.0
    piu_edc = df_piu[df_piu['method'].str.upper().str.contains('EDC', na=False)]['amount'].sum() if not df_piu.empty and 'amount' in df_piu.columns else 0.0
    piu_qris = df_piu[df_piu['method'].str.upper().str.contains('QRIS', na=False)]['amount'].sum() if not df_piu.empty and 'amount' in df_piu.columns else 0.0
    piu_va = df_piu[df_piu['method'].str.upper().str.contains('VIRTUAL', na=False) | df_piu['method'].str.upper().str.contains('VA', na=False)]['amount'].sum() if not df_piu.empty and 'amount' in df_piu.columns else 0.0

    tot_tunai_sys = (df_tx['pay_tunai'].sum() if not df_tx.empty and 'pay_tunai' in df_tx.columns else 0) + piu_tunai
    tot_transfer_sys = (df_tx['pay_transfer'].sum() if not df_tx.empty and 'pay_transfer' in df_tx.columns else 0) + piu_transfer
    tot_edc_sys = (df_tx['pay_edc'].sum() if not df_tx.empty and 'pay_edc' in df_tx.columns else 0) + piu_edc
    tot_qris_sys = (df_tx['pay_qris'].sum() if not df_tx.empty and 'pay_qris' in df_tx.columns else 0) + piu_qris
    tot_va_sys = (df_tx['pay_va'].sum() if not df_tx.empty and 'pay_va' in df_tx.columns else 0) + piu_va
    tot_pengakuan_sys = df_tx['pay_pengakuan_bendahara'].sum() if not df_tx.empty and 'pay_pengakuan_bendahara' in df_tx.columns else 0
    grand_kanal_sys = tot_tunai_sys + tot_transfer_sys + tot_edc_sys + tot_qris_sys + tot_va_sys + tot_pengakuan_sys

    # --- METRICS DASHBOARD ---
    tot_diskon = df_items['discount'].sum() if not df_items.empty and 'discount' in df_items.columns else 0
    tot_subtotal_kotor = df_items['subtotal'].sum() if not df_items.empty else 0
    tot_qty_item = int(df_items['qty'].sum()) if not df_items.empty and 'qty' in df_items.columns else 0
    tot_piutang_masuk = df_piu['amount'].sum() if not df_piu.empty and 'amount' in df_piu.columns else 0
    
    m1, m2, m3, m4 = st.columns(4)
    m1.markdown(f"<div class='metric-card'>Total Pendapatan Netto<br><b>{format_rupiah(tot_subtotal_kotor)}</b></div>", unsafe_allow_html=True)
    m2.markdown(f"<div class='metric-card'>Total Item Terjual (Qty)<br><b>{tot_qty_item:,} Item</b></div>", unsafe_allow_html=True)
    m3.markdown(f"<div class='metric-card'>Total Diskon Karyawan<br><b>{format_rupiah(tot_diskon)}</b></div>", unsafe_allow_html=True)
    m4.markdown(f"<div class='metric-card'>Total Pelunasan Piutang<br><b>{format_rupiah(tot_piutang_masuk)}</b></div>", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    # --- FITUR PENCOCOKAN REKENING KORAN ---
    st.markdown('<div class="bank-match-box">', unsafe_allow_html=True)
    st.markdown("##### 🏦 Fitur Pencocokan Rekening Koran / Mutasi Bank Fisik", unsafe_allow_html=True)
    bm1, bm2 = st.columns([2, 2])
    with bm1:
        bank_actual_input = st.number_input("Masukkan Total Mutasi Masuk / Rekening Koran (Rp)", min_value=0.0, value=float(grand_kanal_sys), step=10000.0, format="%.0f")
    with bm2:
        selisih_bank = bank_actual_input - grand_kanal_sys
        warna_selisih = "#10B981" if selisih_bank == 0 else "#EF4444"
        st.markdown(f"<div style='padding-top:15px;'>Total Sistem Kasir: <b>{format_rupiah(grand_kanal_sys)}</b><br>Selisih (Variance): <span style='color:{warna_selisih}; font-weight:800; font-size:16px;'>{format_rupiah(selisih_bank)}</span></div>", unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

    # --- TAMPILAN TABEL HIERARKI & FITUR CETAK ---
    tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
        "📊 Detail Item", 
        "📂 Rekap per Unit", 
        "🩺 Rekap per Tindakan", 
        "📈 Grafik Unit", 
        "💳 Kanal Pembayaran", 
        "📑 Penerimaan Piutang"
    ])

    with tab1:
        c_prt1, _ = st.columns([2, 8])
        with c_prt1:
            if st.button("🖨️ Cetak Laporan Bendahara", use_container_width=True):
                components.html("<script>window.print();</script>", height=0)

        st.markdown("##### 📈 Rincian Lengkap Berdasarkan Kategori Layanan, Unit, Tindakan, dan No. Kertas")
        if not df_items.empty:
            html_remun = "<table class='tbl-report'>"
            html_remun += "<tr><th>No</th><th>No. Kertas / Ref</th><th>Shift</th><th>Kasir</th><th>Kategori Layanan Utama</th><th>Unit Layanan</th><th>Nama Tindakan Spesifik</th><th style='text-align:center;'>Qty</th><th style='text-align:right;'>Diskon (Rp)</th><th style='text-align:right;'>Subtotal Netto (Rp)</th></tr>"
            
            if 'discount' not in df_items.columns:
                df_items['discount'] = 0.0

            df_sorted = df_items.sort_values(by=['service_category_name', 'category_name', 'action_name', 'receipt_no'])
            
            grand_tot_sub = 0
            grand_tot_disc = 0
            grand_tot_qty = 0
            
            current_scat = ""
            row_idx = 1

            for _, r in df_sorted.iterrows():
                scat_name = r['service_category_name']
                
                if scat_name != current_scat:
                    current_scat = scat_name
                    html_remun += f"<tr class='subtotal-unit-row'><td colspan='10'>🏷️ KATEGORI LAYANAN UTAMA: <b>{current_scat}</b></td></tr>"

                grand_tot_qty += int(r['qty'])
                grand_tot_disc += float(r['discount'])
                grand_tot_sub += float(r['subtotal'])
                
                ksr_name = str(r.get('cashier_username', '-')).upper()
                shf_name = str(r.get('shift', '-'))
                html_remun += f"<tr><td>{row_idx}</td><td><b>#{r['receipt_no']}</b></td><td>{shf_name}</td><td>{ksr_name}</td><td>{r['service_category_name']}</td><td style='padding-left: 10px;'>↳ {r['category_name']}</td><td><b>{r['action_name']}</b></td><td style='text-align:center;'>{r['qty']}</td><td style='text-align:right; color:#D97706;'>{format_angka(r['discount'])}</td><td style='text-align:right;'>{format_angka(r['subtotal'])}</td></tr>"
                row_idx += 1
            
            html_remun += f"<tr class='total-row'><td colspan='7'>TOTAL KESELURUHAN</td><td style='text-align:center;'>{grand_tot_qty}</td><td style='text-align:right; color:#D97706;'>{format_angka(grand_tot_disc)}</td><td style='text-align:right;'>{format_rupiah(grand_tot_sub).replace('Rp ', '')}</td></tr>"
            html_remun += "</table>"
            st.markdown(html_remun, unsafe_allow_html=True)
        else:
            st.info("Tidak ada data tindakan yang sesuai dengan filter atau pencarian Anda.")

    with tab2:
        st.markdown("##### 📂 Ringkasan Totalan Pendapatan Per Unit Layanan & Kategori Utama")
        if not df_items.empty:
            df_unit_summary = df_items.groupby(['service_category_name', 'category_name']).agg({'qty': 'sum', 'discount': 'sum', 'subtotal': 'sum'}).reset_index()
            
            html_us = "<table class='tbl-report'>"
            html_us += "<tr><th>No</th><th>Kategori Layanan Utama</th><th>Unit Layanan</th><th style='text-align:center;'>Total Qty</th><th style='text-align:right;'>Total Diskon (Rp)</th><th style='text-align:right;'>Total Netto (Rp)</th></tr>"
            
            u_qty = 0
            u_disc = 0.0
            u_sub = 0.0
            
            for idx, r in df_unit_summary.iterrows():
                u_qty += int(r['qty'])
                u_disc += float(r['discount'])
                u_sub += float(r['subtotal'])
                html_us += f"<tr><td>{idx+1}</td><td>{r['service_category_name']}</td><td><b>{r['category_name']}</b></td><td style='text-align:center;'>{r['qty']}</td><td style='text-align:right; color:#D97706;'>{format_angka(r['discount'])}</td><td style='text-align:right;'>{format_angka(r['subtotal'])}</td></tr>"
            
            html_us += f"<tr class='total-row'><td colspan='3'>TOTAL KESELURUHAN UNIT</td><td style='text-align:center;'>{u_qty}</td><td style='text-align:right; color:#D97706;'>{format_angka(u_disc)}</td><td style='text-align:right;'>{format_angka(u_sub)}</td></tr>"
            html_us += "</table>"
            st.markdown(html_us, unsafe_allow_html=True)
        else:
            st.info("Tidak ada data unit layanan pada periode ini.")

    with tab3:
        st.markdown("##### 🩺 Rekapitulasi Totalan Berdasarkan Masing-Masing Tindakan Spesifik")
        if not df_items.empty:
            df_action_summary = df_items.groupby(['service_category_name', 'category_name', 'action_name']).agg({'qty': 'sum', 'discount': 'sum', 'subtotal': 'sum'}).reset_index()
            df_action_summary = df_action_summary.sort_values(by=['service_category_name', 'category_name', 'action_name'])

            html_as = "<table class='tbl-report'>"
            html_as += "<tr><th>No</th><th>Kategori Layanan</th><th>Unit Layanan</th><th>Nama Tindakan Spesifik</th><th style='text-align:center;'>Total Qty</th><th style='text-align:right;'>Total Diskon (Rp)</th><th style='text-align:right;'>Total Netto (Rp)</th></tr>"
            
            a_qty = 0
            a_disc = 0.0
            a_sub = 0.0
            
            for idx, r in df_action_summary.iterrows():
                a_qty += int(r['qty'])
                a_disc += float(r['discount'])
                a_sub += float(r['subtotal'])
                html_as += f"<tr><td>{idx+1}</td><td>{r['service_category_name']}</td><td>{r['category_name']}</td><td><b>{r['action_name']}</b></td><td style='text-align:center;'>{r['qty']}</td><td style='text-align:right; color:#D97706;'>{format_angka(r['discount'])}</td><td style='text-align:right;'>{format_angka(r['subtotal'])}</td></tr>"
            
            html_as += f"<tr class='total-row'><td colspan='4'>TOTAL KESELURUHAN TINDAKAN</td><td style='text-align:center;'>{a_qty}</td><td style='text-align:right; color:#D97706;'>{format_angka(a_disc)}</td><td style='text-align:right;'>{format_angka(a_sub)}</td></tr>"
            html_as += "</table>"
            st.markdown(html_as, unsafe_allow_html=True)
        else:
            st.info("Tidak ada data tindakan spesifik pada periode ini.")

    with tab4:
        st.markdown("##### 📈 Grafik Visualisasi Pendapatan Berdasarkan Kategori & Unit Layanan")
        if not df_items.empty:
            df_chart = df_items.groupby('category_name')['subtotal'].sum().reset_index()
            df_chart = df_chart.set_index('category_name')
            st.bar_chart(df_chart)
        else:
            st.info("Tidak ada data untuk ditampilkan pada grafik.")

    with tab5:
        st.markdown("##### Rekapitulasi Metode Pembayaran (Transaksi + Pelunasan Piutang)")
        if not df_tx.empty or not df_piu.empty:
            html_pm = "<table class='tbl-report'>"
            html_pm += "<tr><th>Metode Pembayaran / Kanal</th><th style='text-align:right;'>Nominal (Rp)</th></tr>"
            html_pm += f"<tr><td>Tunai (Transaksi + Piutang)</td><td style='text-align:right;'>{format_angka(tot_tunai_sys)}</td></tr>"
            html_pm += f"<tr><td>Transfer Bank (Transaksi + Piutang)</td><td style='text-align:right;'>{format_angka(tot_transfer_sys)}</td></tr>"
            html_pm += f"<tr><td>EDC Kartu (Transaksi + Piutang)</td><td style='text-align:right;'>{format_angka(tot_edc_sys)}</td></tr>"
            html_pm += f"<tr><td>QRIS (Transaksi + Piutang)</td><td style='text-align:right;'>{format_angka(tot_qris_sys)}</td></tr>"
            html_pm += f"<tr><td>Virtual Account / VA (Transaksi + Piutang)</td><td style='text-align:right;'>{format_angka(tot_va_sys)}</td></tr>"
            html_pm += f"<tr><td>Pengakuan / Potong Uang Muka</td><td style='text-align:right;'>{format_angka(tot_pengakuan_sys)}</td></tr>"
            html_pm += f"<tr class='total-row'><td>Total Seluruh Kanal Pembayaran</td><td style='text-align:right;'>{format_angka(grand_kanal_sys)}</td></tr>"
            html_pm += "</table>"
            st.markdown(html_pm, unsafe_allow_html=True)
        else:
            st.info("Tidak ada data transaksi pembayaran.")

    with tab6:
        st.markdown("##### Rincian Pelunasan / Penerimaan Piutang Pasien")
        if not df_piu.empty:
            html_piu_rekon = "<table class='tbl-report'>"
            html_piu_rekon += "<tr><th>No</th><th>Tanggal Bayar</th><th>Shift</th><th>Metode Bayar</th><th>Petugas / Kasir</th><th style='text-align:right;'>Nominal (Rp)</th></tr>"
            tot_p_rekon = 0.0
            for idx, r in df_piu.iterrows():
                amt = float(r['amount'] or 0)
                tot_p_rekon += amt
                html_piu_rekon += f"<tr><td>{idx+1}</td><td>{r.get('pay_date', '-')}</td><td>{r.get('shift', '-')}</td><td>{r.get('method', '-')}</td><td>{str(r.get('input_by', '-')).upper()}</td><td style='text-align:right; color:#10B981; font-weight:700;'>{format_angka(amt)}</td></tr>"
            html_piu_rekon += f"<tr class='total-row'><td colspan='5'>Total Penerimaan Piutang</td><td style='text-align:right; color:#10B981;'>{format_angka(tot_p_rekon)}</td></tr>"
            html_piu_rekon += "</table>"
            st.markdown(html_piu_rekon, unsafe_allow_html=True)
        else:
            st.info("Tidak ada data penerimaan piutang pada periode ini.")

    # --- TOMBOL EKSPOR EXCEL MULTI-SHEET ---
    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("📥 Unduh Laporan Rekonsiliasi Lengkap (Excel)", use_container_width=True, type="primary"):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            if not df_items.empty:
                df_items.to_excel(writer, sheet_name='Detail Item & No Kertas', index=False)
                df_unit_summary.to_excel(writer, sheet_name='Rekap Per Unit', index=False)
                df_action_summary.to_excel(writer, sheet_name='Rekap Per Tindakan', index=False)
            if not df_tx.empty:
                df_tx.to_excel(writer, sheet_name='Transaksi Utama', index=False)
            if not df_piu.empty:
                df_piu.to_excel(writer, sheet_name='Penerimaan Piutang', index=False)
        
        st.download_button(
            label="Klik Disini untuk Mengunduh File Excel",
            data=output.getvalue(),
            file_name=f"Rekon_Bendahara_{tipe_laporan}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )